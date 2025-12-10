import asyncio
import os
import re
from typing import Optional, Dict, List

import asyncpg
import docx
import pandas as pd
from PyPDF2 import PdfReader
from dotenv import load_dotenv
from superagentx.handler.base import BaseHandler
from superagentx.handler.decorators import tool
from superagentx.llm.openai import AzureOpenAI

load_dotenv()

DB_URL = os.getenv("DB_URL")
AZURE_OPENAI_API_KEY = os.getenv("AZURE_OPENAI_API_KEY")
AZURE_OPENAI_ENDPOINT = os.getenv("AZURE_API_BASE")
RAG_FOLDER_PATH = os.getenv("RAG_FOLDER_PATH")
OPENAI_API_VERSION = os.getenv("OPENAI_API_VERSION")


class RAGIngestHandler(BaseHandler):

    def __init__(
            self,
            db_url: str = DB_URL,
            embedding_dim: int = 1536,
            preview_chars: int = 10000,
            chunk_base_size: int = 1500,
            chunk_overlap: int = 300,
            concurrency: int = 10,
    ):
        super().__init__()
        self.db_url = db_url
        self.embedding_dim = embedding_dim
        self.chunk_base_size = chunk_base_size
        self.chunk_overlap = chunk_overlap
        self.preview_chars = preview_chars
        self.concurrency = concurrency

        self.azure_client = AzureOpenAI(
            api_key=AZURE_OPENAI_API_KEY,
            api_version=OPENAI_API_VERSION,
            azure_endpoint=AZURE_OPENAI_ENDPOINT,
        )

        self.pool: Optional[asyncpg.pool.Pool] = None
        self._sem = asyncio.Semaphore(self.concurrency)

    async def init_db(self):
        if self.pool:
            return

        self.pool = await asyncpg.create_pool(self.db_url, min_size=1, max_size=5)

        async with self.pool.acquire() as conn:
            await conn.execute(
                f"""
                CREATE EXTENSION IF NOT EXISTS vector;

                CREATE TABLE IF NOT EXISTS documents (
                    id SERIAL PRIMARY KEY,
                    filename TEXT UNIQUE,
                    full_text TEXT,
                    extension TEXT,
                    embedding vector({self.embedding_dim})
                );

                CREATE TABLE IF NOT EXISTS chunks (
                    id SERIAL PRIMARY KEY,
                    document_id INT REFERENCES documents(id) ON DELETE CASCADE,
                    chunk_text TEXT,
                    embedding vector({self.embedding_dim})
                );

                CREATE TABLE IF NOT EXISTS table_rows (
                    id SERIAL PRIMARY KEY,
                    document_id INT REFERENCES documents(id) ON DELETE CASCADE,
                    sheet_name TEXT,
                    row_text TEXT
                );
                """
            )

    def safe_sql_name(self, name: str) -> str:
        name = (name or "").strip().lower()
        name = name.replace(" ", "_")
        name = re.sub(r"[^a-z0-9_]", "", name)
        if not name:
            name = "t"
        if re.match(r"^[0-9]", name):
            name = "t_" + name
        return name

    def emb_sync(self, text: str):
        r = self.azure_client.embeddings.create(
            model="text-embedding-ada-002",
            input=text
        )
        return r.data[0].embedding

    async def emb(self, text: str):
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(None, self.emb_sync, text)

    def to_vec(self, arr):
        return "[" + ",".join(map(str, arr)) + "]"


    def load_document(self, path: str):
        supported = [".pdf", ".docx", ".csv", ".xlsx", ".xls", ".txt", ".json"]

        if os.path.isdir(path):
            out = []
            for root, _, files in os.walk(path):
                for f in files:
                    if os.path.splitext(f)[1].lower() in supported:
                        out.append(os.path.join(root, f))
            return out

        if os.path.isfile(path):
            return [path]

        raise ValueError("Invalid path.")

    def stream_pdf(self, p):
        reader = PdfReader(p)
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                yield t

    def stream_docx(self, p):
        doc = docx.Document(p)
        for para in doc.paragraphs:
            if para.text.strip():
                yield para.text

    def stream_txt(self, p):
        with open(p, "r", encoding="utf-8", errors="ignore") as f:
            for line in f:
                if line.strip():
                    yield line.strip()

    def stream_csv(self, p):
        try:
            for chunk in pd.read_csv(p, chunksize=1, dtype=str):
                row = chunk.iloc[0].fillna("")
                yield "\n".join(f"{c}: {row[c]}" for c in chunk.columns)
        except Exception:
            yield from self.stream_txt(p)

    def stream_xlsx(self, p):
        try:
            all_sheets = pd.read_excel(p, engine="openpyxl", dtype=str, sheet_name=None)
            for sheet_name, df in all_sheets.items():
                df = df.fillna("")
                for _, row in df.iterrows():
                    row_text = "\n".join(f"{c}: {row[c]}" for c in df.columns)
                    yield sheet_name, row_text
        except Exception as e:
            yield None, f"Unable to load {p}: {str(e)}"


    def dynamic_stream_chunker(self, generator, base, overlap):
        buf = ""
        for piece in generator:
            piece = piece.strip()
            if not piece:
                continue

            if buf:
                buf += "\n" + piece
            else:
                buf = piece

            if len(buf) >= base:
                yield buf
                buf = buf[-overlap:]

        if buf.strip():
            yield buf

    async def _ensure_table_for_file(self, conn: asyncpg.Connection, table_name: str, columns: List[str]):
        await conn.execute(
            f"""
            CREATE TABLE IF NOT EXISTS "{table_name}" (
                id SERIAL PRIMARY KEY,
                sheet_name TEXT
            );
            """
        )

        rows = await conn.fetch(
            """
            SELECT column_name FROM information_schema.columns
            WHERE table_name = $1
            """,
            table_name.lower()
        )
        existing = {r["column_name"] for r in rows}
        for col in columns:
            if col not in existing:
                await conn.execute(f'ALTER TABLE "{table_name}" ADD COLUMN "{col}" TEXT;')

    async def process_file(self, file_path: str):
        await self.init_db()
        filename = os.path.basename(file_path)
        ext = os.path.splitext(filename)[1].lower()

        if ext == ".pdf":
            gen = self.stream_pdf(file_path)
            store_as_table = False
        elif ext == ".txt":
            gen = self.stream_txt(file_path)
            store_as_table = False
        elif ext in [".csv", ".xlsx", ".xls"]:
            store_as_table = True
            gen = None
        elif ext == ".docx":
            gen = self.stream_docx(file_path)
            store_as_table = False
        else:
            return {"file": filename, "error": f"Unsupported file type: {ext}"}

        async with self.pool.acquire() as conn:
            doc_id = await conn.fetchval(
                "INSERT INTO documents(filename, full_text, extension, embedding) "
                "VALUES ($1,$2,$3,$4::vector) "
                "ON CONFLICT(filename) DO UPDATE SET full_text=EXCLUDED.full_text, extension=EXCLUDED.extension RETURNING id;",
                filename, "", ext, self.to_vec([0.0] * self.embedding_dim)
            )

        if store_as_table and ext in [".csv", ".xlsx", ".xls"]:
            if ext == ".csv":
                try:
                    all_sheets: Dict[str, pd.DataFrame] = {"sheet1": pd.read_csv(file_path, dtype=str).fillna("")}
                except Exception as e:
                    return {"file": filename, "error": f"Failed reading CSV: {str(e)}"}
            else:
                try:
                    all_sheets = pd.read_excel(file_path, sheet_name=None, engine="openpyxl", dtype=str)
                    all_sheets = {k: v.fillna("") for k, v in all_sheets.items()}
                except Exception as e:
                    return {"file": filename, "error": f"Failed reading XLSX: {str(e)}"}

            base_name = os.path.splitext(filename)[0]
            table_name = self.safe_sql_name(f"{base_name}_{ext.replace('.', '')}")

            union_cols = []
            for sheet_df in all_sheets.values():
                for c in list(sheet_df.columns):
                    safe = self.safe_sql_name(str(c))
                    if safe not in union_cols:
                        union_cols.append(safe)
            columns = union_cols

            async with self.pool.acquire() as conn:
                await self._ensure_table_for_file(conn, table_name, columns)

            insert_count = 0
            async with self.pool.acquire() as conn:
                all_insert_cols = ["sheet_name"] + columns
                placeholders = ", ".join(f"${i+1}" for i in range(len(all_insert_cols)))
                quoted_cols = ", ".join(f'"{c}"' for c in all_insert_cols)
                insert_sql = f'INSERT INTO "{table_name}" ({quoted_cols}) VALUES ({placeholders});'

                for sheet_name, df in all_sheets.items():
                    for _, row in df.iterrows():
                        row_map = {
                            self.safe_sql_name(str(col)): ("" if pd.isna(row[col]) else str(row[col]))
                            for col in df.columns
                        }
                        vals = [sheet_name] + [row_map.get(c, "") for c in columns]

                        await conn.execute(insert_sql, *vals)
                        insert_count += 1
                        row_text = "\n".join(f"{col}: {row_map.get(self.safe_sql_name(str(col)), '')}" for col in df.columns)
                        await conn.execute(
                            """
                            INSERT INTO table_rows(document_id, sheet_name, row_text)
                            VALUES ($1, $2, $3)
                            """,
                            doc_id, sheet_name, row_text
                        )

            return {"file": filename, "status": "stored_as_table", "table": table_name, "rows_inserted": insert_count}

        preview_parts, preview_len = [], 0
        tasks = []

        async def handle_chunk(chunk_text):
            emb = await self.emb(chunk_text)
            vec = self.to_vec(emb)
            async with self.pool.acquire() as c:
                await c.execute(
                    "INSERT INTO chunks(document_id, chunk_text, embedding) VALUES ($1,$2,$3::vector)",
                    doc_id, chunk_text, vec
                )

        async def handle_row(row_text, sheet_name=None):
            async with self.pool.acquire() as c:
                await c.execute(
                    "INSERT INTO table_rows(document_id, sheet_name, row_text) VALUES ($1,$2,$3)",
                    doc_id, sheet_name, row_text
                )

        if gen is None:
            return {"file": filename, "status": "stored_as_table_no_chunks"}

        for piece in self.dynamic_stream_chunker(gen, self.chunk_base_size, self.chunk_overlap):
            if preview_len < self.preview_chars:
                need = self.preview_chars - preview_len
                preview_parts.append(piece[:need])
                preview_len += len(piece[:need])
            tasks.append(asyncio.create_task(handle_chunk(piece)))

        await asyncio.gather(*tasks)

        preview_text = "".join(preview_parts)
        try:
            p_emb = await self.emb(preview_text or " ")
        except Exception:
            p_emb = [0.0] * self.embedding_dim

        async with self.pool.acquire() as conn:
            await conn.execute(
                "UPDATE documents SET full_text=$1, embedding=$2::vector WHERE id=$3",
                preview_text, self.to_vec(p_emb), doc_id
            )

        return {"file": filename, "status": "ingested"}

    @tool
    async def ingest(self, path: Optional[str] = None):
        await self.init_db()
        if path is None:
            path = RAG_FOLDER_PATH
        files = self.load_document(path)
        results = []
        for f in files:
            try:
                r = await self.process_file(f)
            except Exception as e:
                r = {"file": os.path.basename(f), "error": str(e)}
            results.append(r)
        return {"processed": len(files), "details": results}